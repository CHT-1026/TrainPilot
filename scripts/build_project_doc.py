"""Build the Word companion from the authoritative Markdown project plan.

Requires python-docx and Pillow. Rendering and visual QA are separate steps.
Run from any directory: python scripts/build_project_doc.py
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "product-plan.md"
OUTPUT = ROOT / "docs" / "TrainPilot项目产品与实施方案.docx"
ASSET = ROOT / "docs" / "assets" / "architecture.png"


def draw_architecture():
    ASSET.parent.mkdir(parents=True, exist_ok=True)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if not font_path.exists():
        font_path = Path("C:/Windows/Fonts/simsun.ttc")
    if not font_path.exists():
        raise RuntimeError("Set an available CJK font path for the architecture image")
    large = ImageFont.truetype(str(font_path), 31)
    small = ImageFont.truetype(str(font_path), 25)
    im = Image.new("RGB", (1440, 440), "white")
    d = ImageDraw.Draw(im)
    boxes = [(45, 130, 440, 320), (525, 130, 920, 320), (1005, 130, 1400, 320)]
    text = [
        ("Agent 端", "模型 API 与状态编排", "读取证据  选择工具"),
        ("实验后端", "接口 校验 SQLite Worker", "任务与预算控制"),
        ("模型训练端", "PyTorch 与小语言模型", "训练 指标 检查点"),
    ]
    for box, lines in zip(boxes, text):
        d.rounded_rectangle(box, radius=12, fill="#F1F5F9", outline="#64748B", width=2)
        center = (box[0] + box[2]) // 2
        for y, line, font in zip((162, 221, 270), lines, (large, small, small)):
            d.text((center, y), line, font=font, fill="#111827", anchor="mm")
    for x in (440, 920):
        d.line((x + 7, 201, x + 76, 201), fill="#334155", width=3)
        d.polygon([(x + 76, 201), (x + 65, 194), (x + 65, 208)], fill="#334155")
        d.line((x + 76, 251, x + 7, 251), fill="#334155", width=3)
        d.polygon([(x + 7, 251), (x + 18, 244), (x + 18, 258)], fill="#334155")
    d.text((720, 67), "用户任务 → 检查与诊断 → 受控实验 → 结果比较", font=large, fill="#111827", anchor="mm")
    d.text((720, 386), "训练的小模型是实验对象   Agent 使用现成模型 API", font=small, fill="#334155", anchor="mm")
    im.save(ASSET)


def set_font(style, size):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def rich_text(p, text):
    pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*")
    start = 0
    for m in pattern.finditer(text):
        p.add_run(text[start:m.start()])
        if m.group(1):
            link = OxmlElement("w:hyperlink")
            rid = p.part.relate_to(m.group(2), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rid)
            r = OxmlElement("w:r")
            pr = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "174A70")
            pr.append(color)
            r.append(pr)
            t = OxmlElement("w:t")
            t.text = m.group(1)
            r.append(t)
            link.append(r)
            p._p.append(link)
        else:
            run = p.add_run(m.group(3) or m.group(4))
            if m.group(4):
                run.bold = True
        start = m.end()
    p.add_run(text[start:])


def add_table(doc, lines):
    rows = [[v.strip() for v in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", c.replace(" ", "")) for c in row)]
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.autofit = False
    widths = [4.0, 6.0, 7.0] if cols == 3 else [6.0, 11.0]
    if rows[0][0] in {"实验", "维度"}:
        widths = [3.2, 7.0, 6.8]
    if rows[0][0] == "阶段":
        widths = [3.2, 2.5, 11.3]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, value in enumerate(row):
            cell = cells[ci]
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                for key, val in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    border.set(qn(f"w:{key}"), val)
                borders.append(border)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "75" if edge in ("top", "bottom") else "110")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            tcpr.append(margins)
            if idx == 0 or idx % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EEF4" if idx == 0 else "F6F8FA")
                tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            rich_text(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.bold = idx == 0
        if idx == 0:
            header = OxmlElement("w:tblHeader")
            table.rows[-1]._tr.get_or_add_trPr().append(header)
        no_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(no_split)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.4
    spacer.add_run().font.size = Pt(2)


def build():
    draw_architecture()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.7)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    for name, size in (("Normal", 10.5), ("Title", 22), ("Heading 1", 17), ("Heading 2", 12), ("List Bullet", 10)):
        set_font(doc.styles[name], size)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(5), 1.12
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(9), Pt(6)
        pf.keep_with_next = True
    doc.core_properties.title = "TrainPilot 项目产品与实施方案"
    doc.core_properties.subject = "Agent 端与模型训练端的技术设计和学习实施路线"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "TrainPilot, AI Infra, Agent, PyTorch"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("TrainPilot  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
        elif line.startswith("```mermaid"):
            while i < len(lines) and lines[i] != "```":
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ASSET), width=Cm(17))
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        else:
            is_bullet = line.startswith("- ")
            p = doc.add_paragraph(style="List Bullet" if is_bullet else "Normal")
            rich_text(p, line[2:] if is_bullet else line)
        i += 1
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print("Designed sections: 8. Actual page count must be checked by rendering.")


if __name__ == "__main__":
    build()

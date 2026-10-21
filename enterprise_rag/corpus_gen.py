"""企业语料生成 CLI：135 篇文档落盘为 md/docx/pdf/csv/xlsx，并输出 manifest.jsonl。

用法（backend 目录）：
    .venv312\\Scripts\\python.exe -m enterprise_rag.corpus_gen
"""
import csv
import hashlib
import json
import shutil
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from enterprise_rag import config
from enterprise_rag.corpus_specs2 import build_documents

CORPUS_ROOT = config.PACKAGE_ROOT / "corpus"
META_FIELDS = [
    ("所属部门", "department"),
    ("作者", "author"),
    ("创建时间", "created"),
    ("更新时间", "updated"),
    ("版本号", "version"),
    ("文档类型", "category"),
    ("文档族", "family"),
    ("生效日期", "effective"),
]


def _meta_lines(document: dict) -> list[str]:
    return [f"{label}：{document[key]}" for label, key in META_FIELDS]


def _table_markdown(table: dict) -> str:
    headers = table["headers"]
    lines = ["| " + " | ".join(str(item) for item in headers) + " |"]
    lines.append("| " + " | ".join("---" for _ in headers) + " |")
    for row in table["rows"]:
        lines.append("| " + " | ".join(str(item) for item in row) + " |")
    return "\n".join(lines)


def write_markdown(document: dict, path: Path) -> None:
    lines = [f"# {document['title']}", ""]
    lines.extend(_meta_lines(document))
    lines.append("")
    for heading, body in document["sections"]:
        lines.append(f"## {heading}")
        lines.append("")
        lines.append(body)
        lines.append("")
    for table in document["tables"]:
        lines.append(f"## {table['name']}")
        lines.append("")
        lines.append(_table_markdown(table))
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_docx(document: dict, path: Path) -> None:
    import docx
    from docx.shared import Pt

    document_obj = docx.Document()
    document_obj.core_properties.title = document["title"]
    document_obj.core_properties.author = document["author"]
    document_obj.core_properties.subject = f"{document['department']} · {document['category']}"
    document_obj.core_properties.comments = f"版本 {document['version']}"
    document_obj.core_properties.category = document["category"]
    document_obj.core_properties.created = datetime.fromisoformat(document["created"])
    document_obj.core_properties.modified = datetime.fromisoformat(document["updated"])

    document_obj.add_heading(document["title"], level=0)
    for line in _meta_lines(document):
        paragraph = document_obj.add_paragraph(line)
        paragraph.runs[0].font.size = Pt(9)
    for heading, body in document["sections"]:
        document_obj.add_heading(heading, level=1)
        document_obj.add_paragraph(body)
    for table in document["tables"]:
        document_obj.add_heading(table["name"], level=1)
        table_obj = document_obj.add_table(rows=1, cols=len(table["headers"]))
        table_obj.style = "Table Grid"
        for index, header in enumerate(table["headers"]):
            table_obj.rows[0].cells[index].text = str(header)
        for row in table["rows"]:
            cells = table_obj.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)
    document_obj.save(str(path))


def _pdf_styles():
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass
    font = "STSong-Light"
    return {
        "title": ParagraphStyle("title", fontName=font, fontSize=18, leading=26, spaceAfter=10),
        "meta": ParagraphStyle("meta", fontName=font, fontSize=8.5, leading=13, textColor="#555555"),
        "heading": ParagraphStyle("heading", fontName=font, fontSize=12.5, leading=20, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17, spaceAfter=6),
    }


def write_pdf(document: dict, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = _pdf_styles()
    frame_width = A4[0] - 40 * mm
    story = [Paragraph(escape(document["title"]), styles["title"])]
    for line in _meta_lines(document):
        story.append(Paragraph(escape(line), styles["meta"]))
    story.append(Spacer(1, 6 * mm))
    for heading, body in document["sections"]:
        story.append(Paragraph(escape(heading), styles["heading"]))
        story.append(Paragraph(escape(body), styles["body"]))
    for table in document["tables"]:
        story.append(Paragraph(escape(table["name"]), styles["heading"]))
        headers = [Paragraph(escape(str(item)), styles["meta"]) for item in table["headers"]]
        rows = [headers]
        for row in table["rows"]:
            rows.append([Paragraph(escape(str(item)), styles["meta"]) for item in row])
        column_width = frame_width / len(table["headers"])
        table_obj = Table(rows, colWidths=[column_width] * len(table["headers"]), repeatRows=1)
        table_obj.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4D6")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(table_obj)
        story.append(Spacer(1, 4 * mm))
    pdf = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=document["title"],
        author=document["author"],
    )
    pdf.build(story)


def write_csv(document: dict, path: Path) -> None:
    table = document["tables"][0]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(table["headers"])
        writer.writerows(table["rows"])


def write_xlsx(document: dict, path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    info = workbook.active
    info.title = "文档信息"
    info.append(["字段", "值"])
    info.append(["标题", document["title"]])
    for label, key in META_FIELDS:
        info.append([label, document[key]])
    for cell in info["A"]:
        cell.font = openpyxl.styles.Font(bold=True)

    for table in document["tables"]:
        sheet = workbook.create_sheet(table["name"][:31])
        sheet.append(table["headers"])
        for row in table["rows"]:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) * 1.6 + 4
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 42)
    workbook.save(str(path))


WRITERS = {
    "md": write_markdown,
    "docx": write_docx,
    "pdf": write_pdf,
    "csv": write_csv,
    "xlsx": write_xlsx,
}


def generate(clean: bool = True) -> dict:
    documents = build_documents()
    ids = [item["doc_id"] for item in documents]
    if len(ids) != len(set(ids)):
        duplicates = sorted({item for item in ids if ids.count(item) > 1})
        raise RuntimeError(f"文档 ID 重复：{duplicates}")

    docs_dir = CORPUS_ROOT / "docs"
    if clean and docs_dir.exists():
        shutil.rmtree(docs_dir)
    docs_dir.mkdir(parents=True, exist_ok=True)

    manifest_path = CORPUS_ROOT / "manifest.jsonl"
    records = []
    for document in documents:
        folder = docs_dir / document["category"]
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{document['doc_id']}.{document['format']}"
        writer = WRITERS[document["format"]]
        writer(document, path)
        content = path.read_bytes()
        records.append(
            {
                "doc_id": document["doc_id"],
                "title": document["title"],
                "namespace": config.RAG_NAMESPACE,
                "category": document["category"],
                "department": document["department"],
                "author": document["author"],
                "doc_family": document["family"],
                "version": document["version"],
                "created_at": document["created"],
                "updated_at": document["updated"],
                "effective_date": document["effective"],
                "format": document["format"],
                "visibility": document["visibility"],
                "status": "active",
                "path": str(path.relative_to(CORPUS_ROOT)).replace("\\", "/"),
                "checksum": hashlib.sha256(content).hexdigest(),
                "size_bytes": len(content),
            }
        )

    records.sort(key=lambda item: item["doc_id"])
    with manifest_path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")

    summary = {
        "documents": len(records),
        "formats": {
            fmt: sum(1 for record in records if record["format"] == fmt)
            for fmt in sorted({record["format"] for record in records})
        },
        "categories": {
            category: sum(1 for record in records if record["category"] == category)
            for category in sorted({record["category"] for record in records})
        },
        "families_with_versions": len(
            {
                record["doc_family"]
                for record in records
                if sum(1 for other in records if other["doc_family"] == record["doc_family"]) > 1
            }
        ),
        "manifest": str(manifest_path),
    }
    return summary


if __name__ == "__main__":
    print(json.dumps(generate(), ensure_ascii=False, indent=2))

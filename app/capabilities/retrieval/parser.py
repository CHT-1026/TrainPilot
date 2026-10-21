"""文档解析：md/txt/docx/pdf -> 纯文本 + 标题/版本元数据。"""
import hashlib
import re
from datetime import date
from pathlib import Path

VERSION_RE = re.compile(r"(\d{4})\s*版")


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "").strip() for page in reader.pages)


def _extract_title(text: str, fallback: str) -> str:
    for line in text.splitlines():
        candidate = line.strip().lstrip("#").strip()
        if candidate:
            return candidate[:200]
    return fallback[:200]


def _extract_version(text: str) -> str:
    match = VERSION_RE.search(text)
    return match.group(1) if match else ""


def parse_document(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = _read_docx(path)
    elif suffix == ".pdf":
        text = _read_pdf(path)
    else:
        text = _read_text(path)

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    title = _extract_title(text, path.stem)
    version = _extract_version(text)
    effective_date = date(int(version), 1, 1) if version else None
    checksum = hashlib.sha256(text.encode("utf-8")).hexdigest()

    return {
        "doc_id": "doc_" + hashlib.sha1(path.name.encode("utf-8")).hexdigest()[:12],
        "title": title,
        "text": text,
        "doc_type": suffix.lstrip("."),
        "version": version,
        "effective_date": effective_date,
        "checksum": checksum,
        "source_path": str(path),
    }
